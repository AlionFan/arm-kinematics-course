# -*- coding: utf-8 -*-
"""PPTX 预览渲染器：把 21 页渲染成 PNG 拼图，用于无 Office 环境下的目检。"""
import math
import os
import sys

from PIL import Image, ImageDraw, ImageFont
from pptx import Presentation
from pptx.util import Emu
from pptx.oxml.ns import qn

EMU_PX = 9525.0
W, H = 1280, 720
FDIR = 'C:/Windows/Fonts/'
F_CN = FDIR + 'msyh.ttc'
F_CNB = FDIR + 'msyhbd.ttc'
F_EN = FDIR + 'times.ttf'
F_MONO = FDIR + 'consola.ttf'


def px(v):
    return float(v) / EMU_PX


def pick_font(name, bold, size_px):
    if name and 'Consolas' in name:
        p = F_MONO
    elif name and 'Times' in name:
        p = F_EN
    else:
        p = F_CNB if bold else F_CN
    try:
        return ImageFont.truetype(p, int(round(size_px)))
    except Exception:
        return ImageFont.load_default()


def solid_rgb(fill):
    try:
        if fill.type is not None and str(fill.type).startswith('SOLID'):
            c = fill.fore_color.rgb
            return (c[0], c[1], c[2])
    except Exception:
        pass
    return None


def line_rgb(shape):
    try:
        c = shape.line.color.rgb
        return (c[0], c[1], c[2])
    except Exception:
        return None


def line_w(shape):
    try:
        if shape.line.width is not None:
            return max(0.6, px(shape.line.width) * 1.0)
    except Exception:
        pass
    return 1.2


def custgeom_pts(shape, spPr):
    cg = spPr.find(qn('a:custGeom'))
    if cg is None:
        return None
    xfrm = spPr.find(qn('a:xfrm'))
    off = xfrm.find(qn('a:off'))
    ext = xfrm.find(qn('a:ext'))
    ox, oy = int(off.get('x')), int(off.get('y'))
    ex, ey = int(ext.get('cx')), int(ext.get('cy'))
    paths = cg.find(qn('a:pathLst')).findall(qn('a:path'))
    out = []
    for path in paths:
        pw = float(path.get('w') or ex)
        ph = float(path.get('h') or ey)
        sx = ex / pw if pw else 1.0
        sy = ey / ph if ph else 1.0
        for cmd in path:
            tag = cmd.tag.split('}')[1]
            if tag == 'moveTo' or tag == 'lnTo':
                pt = cmd.find(qn('a:pt'))
                x = ox + float(pt.get('x')) * sx
                y = oy + float(pt.get('y')) * sy
                out.append((tag, px(x), px(y)))
    return out


def draw_freeform(d, shape):
    spPr = shape._element.spPr
    pts = custgeom_pts(shape, spPr)
    if not pts:
        return
    col = line_rgb(shape)
    w = line_w(shape)
    if col is None:
        col = (60, 60, 60)
    seq = []
    for tag, x, y in pts:
        if tag == 'moveTo':
            if len(seq) > 1:
                d.line(seq, fill=col, width=max(1, int(round(w))), joint='curve')
            seq = [(x, y)]
        else:
            seq.append((x, y))
    if len(seq) > 1:
        d.line(seq, fill=col, width=max(1, int(round(w))), joint='curve')


def draw_autoshape(d, shape, bg):
    el = shape._element
    spPr = el.find(qn('p:spPr'))
    prst_el = spPr.find(qn('a:prstGeom'))
    prst = prst_el.get('prst') if prst_el is not None else 'rect'
    l, t = px(shape.left), px(shape.top)
    w, h = px(shape.width), px(shape.height)
    box = [l, t, l + w, t + h]
    fill = solid_rgb(shape.fill)
    lc = line_rgb(shape)
    lw = max(1, int(round(line_w(shape))))
    if prst == 'ellipse':
        if fill:
            d.ellipse(box, fill=fill)
        if lc:
            d.ellipse(box, outline=lc, width=lw)
    elif prst == 'donut':
        adj = 25000
        av = spPr.find(qn('a:prstGeom')).find(qn('a:avLst'))
        if av is not None:
            for gd in av.findall(qn('a:gd')):
                if gd.get('name') == 'adj':
                    adj = int(gd.get('fmla').split()[-1])
        ss = min(w, h)
        ring = ss * adj / 100000.0
        if fill:
            d.ellipse(box, fill=fill)
        hole = [l + ring, t + ring, l + w - ring, t + h - ring]
        d.ellipse(hole, fill=bg)
        if lc:
            d.ellipse(box, outline=lc, width=lw)
            d.ellipse(hole, outline=lc, width=lw)
    elif prst == 'roundRect':
        if fill:
            d.rounded_rectangle(box, radius=min(w, h) * 0.28, fill=fill)
        if lc:
            d.rounded_rectangle(box, radius=min(w, h) * 0.28, outline=lc, width=lw)
    else:
        if fill:
            d.rectangle(box, fill=fill)
        if lc:
            d.rectangle(box, outline=lc, width=lw)


def draw_connector(d, shape):
    l, t = px(shape.left), px(shape.top)
    w, h = px(shape.width), px(shape.height)
    el = shape._element
    spPr = el.find(qn('p:spPr'))
    xfrm = spPr.find(qn('a:xfrm'))
    fh = xfrm.get('flipH') == '1'
    fv = xfrm.get('flipV') == '1'
    x1, x2 = (l + w, l) if fh else (l, l + w)
    y1, y2 = (t + h, t) if fv else (t, t + h)
    col = line_rgb(shape) or (60, 60, 60)
    wd = max(1, int(round(line_w(shape))))
    d.line([(x1, y1), (x2, y2)], fill=col, width=wd)
    ln = spPr.find(qn('a:ln'))
    heads = {'headEnd': False, 'tailEnd': False}
    for tag in ('headEnd', 'tailEnd'):
        if ln is not None and ln.find(qn('a:' + tag)) is not None:
            heads[tag] = True
    ang = math.atan2(y2 - y1, x2 - x1)
    size = max(7.0, wd * 3.6)

    def tri(x, y, a):
        for s in (0, 1):
            aa = a + (math.pi - 0.42) * (1 if s else -1)
            d.line([(x, y), (x + size * math.cos(aa), y + size * math.sin(aa))],
                   fill=col, width=wd)
    if heads['tailEnd']:
        tri(x2, y2, ang + math.pi)
    if heads['headEnd']:
        tri(x1, y1, ang)


def draw_text(d, shape):
    tf = shape.text_frame
    l, t = px(shape.left), px(shape.top)
    w, h = px(shape.width), px(shape.height)
    paras = []
    for p in tf.paragraphs:
        runs = []
        for r in p.runs:
            sz = r.font.size.pt if r.font.size else 16
            col = (26, 34, 48)
            try:
                if r.font.color and r.font.color.rgb:
                    c = r.font.color.rgb
                    col = (c[0], c[1], c[2])
            except Exception:
                pass
            runs.append({'t': r.text, 'sz': sz, 'col': col,
                         'b': bool(r.font.bold), 'f': r.font.name})
        ls = p.line_spacing if isinstance(p.line_spacing, float) else 1.2
        al = str(p.alignment) if p.alignment is not None else 'LEFT'
        paras.append({'runs': runs, 'ls': ls, 'al': al})
    if not any(pp['runs'] for pp in paras):
        return
    # 逐段排布
    lines = []  # (x, y, text, font, color)
    cur_y = t
    for pp in paras:
        if not pp['runs']:
            fs = 16
            cur_y += fs * 4 / 3.0 * pp['ls']
            continue
        fs = max(r['sz'] for r in pp['runs'])
        fpx = fs * 4 / 3.0
        fonts = {}
        # 组装 token
        toks = []
        for r in pp['runs']:
            f = pick_font(r['f'], r['b'], r['sz'] * 4 / 3.0)
            fonts.setdefault(r['col'], f)
            for ch in r['t']:
                toks.append((ch, f, r['col']))
        # 贪心折行
        cur = []
        curw = 0.0
        segs = []
        maxw = max(10.0, w - 2)
        for ch, f, col in toks:
            cw = f.getlength(ch)
            if ch == '\n':
                segs.append(cur)
                cur = []
                curw = 0.0
                continue
            if curw + cw > maxw and cur:
                segs.append(cur)
                cur = []
                curw = 0.0
            cur.append((ch, f, col))
            curw += cw
        if cur:
            segs.append(cur)
        lh = fpx * pp['ls'] * 1.02
        for seg in segs:
            sw = sum(f.getlength(c) for c, _f, _c in seg)
            if 'CENTER' in pp['al']:
                sx = l + (w - sw) / 2.0
            elif 'RIGHT' in pp['al']:
                sx = l + w - sw
            else:
                sx = l
            lines.append((sx, cur_y, seg))
            cur_y += lh
    total = cur_y - t
    anchor = str(tf.vertical_anchor)
    dy = 0.0
    if 'MIDDLE' in anchor:
        dy = (h - total) / 2.0
    elif 'BOTTOM' in anchor:
        dy = h - total
    out_lines = []
    for sx, y, seg in lines:
        yy = y + dy
        x = sx
        for ch, f, col in seg:
            out_lines.append((x, yy, ch, f, col))
            x += f.getlength(ch)
    return out_lines, t + dy, total


def render(pptx_path, outdir):
    prs = Presentation(pptx_path)
    os.makedirs(outdir, exist_ok=True)
    paths = []
    for i, s in enumerate(prs.slides, 1):
        bg = (255, 255, 255)
        shapes = list(s.shapes)
        if shapes:
            f = solid_rgb(shapes[0].fill)
            if f:
                bg = f
        img = Image.new('RGB', (W, H), bg)
        # 背景矩形
        d0 = ImageDraw.Draw(img)
        if shapes:
            f = solid_rgb(shapes[0].fill)
            if f:
                d0.rectangle([0, 0, W, H], fill=f)
        texts = []
        for sh in shapes[1:]:
            st = str(sh.shape_type)
            try:
                if 'TEXT_BOX' in st:
                    r = draw_text(d0, sh)
                    if r:
                        texts.append(r)
                elif 'CONNECTOR' in st or 'LINE' in st:
                    draw_connector(d0, sh)
                elif 'FREEFORM' in st:
                    draw_freeform(d0, sh)
                elif 'AUTO_SHAPE' in st or 'SHAPE' in st:
                    draw_autoshape(d0, sh, bg)
                elif 'PICTURE' in st:
                    pass
            except Exception as e:
                print('  warn P%02d %s %s' % (i, st, e))
        for out_lines, top, tot in texts:
            for x, y, ch, f, col in out_lines:
                d0.text((x, y), ch, font=f, fill=col)
        p = os.path.join(outdir, 'p%02d.png' % i)
        img.save(p)
        paths.append(p)
    return paths


def sheet(paths, out, cols=4, cw=480):
    ch = int(cw * H / float(W))
    rows = int(math.ceil(len(paths) / float(cols)))
    canvas = Image.new('RGB', (cols * cw + (cols + 1) * 8, rows * ch + (rows + 1) * 8),
                       (222, 226, 232))
    for i, p in enumerate(paths):
        im = Image.open(p).resize((cw, ch), Image.LANCZOS)
        r, c = divmod(i, cols)
        canvas.paste(im, (8 + c * (cw + 8), 8 + r * (ch + 8)))
    canvas.save(out)
    return out


if __name__ == '__main__':
    src = sys.argv[1] if len(sys.argv) > 1 else '机械臂运动学-前三讲.pptx'
    outd = sys.argv[2] if len(sys.argv) > 2 else 'preview'
    ps = render(src, outd)
    print('rendered', len(ps), 'pages')
    print('sheet ->', sheet(ps, os.path.join(outd, '_sheet.png')))
