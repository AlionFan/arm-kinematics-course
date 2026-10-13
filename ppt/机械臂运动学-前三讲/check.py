# -*- coding: utf-8 -*-
"""版式自检：越界 / 文本溢出 / 压页脚 / 备注完整性。"""
import sys

from pptx import Presentation

from preview import draw_text, px, W, H

FOOTER = {'机械臂运动学 · 前三讲', '%02d / %d'}
src = sys.argv[1] if len(sys.argv) > 1 else '机械臂运动学-前三讲.pptx'
prs = Presentation(src)
issues = []
note_stats = []
for i, s in enumerate(prs.slides, 1):
    txt = ''
    try:
        txt = s.notes_slide.notes_text_frame.text.strip()
    except Exception:
        pass
    note_stats.append((i, len(txt)))
    for sh in s.shapes:
        l, t = px(sh.left), px(sh.top)
        w, h = px(sh.width), px(sh.height)
        r, b = l + w, t + h
        tag = (sh.text_frame.text.strip()[:22] if sh.has_text_frame else '')
        if l < -1 or t < -1 or r > W + 1 or b > H + 1:
            issues.append('P%02d OOB  L%.0f T%.0f R%.0f B%.0f  %r' % (i, l, t, r, b, tag))
        if not sh.has_text_frame:
            continue
        body = sh.text_frame.text.strip()
        if not body or body in ('机械臂运动学 · 前三讲',) or '/' in body and len(body) < 9:
            continue
        res = draw_text(None, sh)
        if res:
            _lines, _top, total = res
            if total > h + 6 and len(body) > 8:
                issues.append('P%02d TEXT 需 %.0f > 框 %.0f  %r' % (i, total, h, body[:34]))
        if b > 645 and t < 648:
            issues.append('P%02d FOOT 压页脚 T%.0f B%.0f  %r' % (i, t, b, tag))

print('slides =', len(prs.slides._sldIdLst))
print('notes  :', ' '.join('P%d=%d字' % (a, b) for a, b in note_stats))
short = [p for p, n in note_stats if n < 60]
print('备注过短页:', short if short else '无')
print('---- 版式问题 ----')
print('\n'.join(issues) if issues else 'OK · 无越界 / 无溢出 / 无压页脚')
